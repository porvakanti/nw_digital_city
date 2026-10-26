"""Building more than one Level 1, and not disturbing the one that exists.

The model began as a single estate and the renderer drew one plate. These
checks hold the generalisation to two promises: an estate that is switched on
is read exactly as Networks is, and switching one on does not change what
Networks produces.

Every fixture is written here rather than read from data/raw, which is never
committed. A test that needs the extract is a test that only runs on the
machine that happens to have it.
"""

from __future__ import annotations

import importlib.util
import pathlib
import sys
import tempfile
import unittest

import openpyxl

REPO = pathlib.Path(__file__).resolve().parent.parent


def load_builder():
    spec = importlib.util.spec_from_file_location(
        "build_city", REPO / "data" / "build_city.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules["build_city"] = module
    spec.loader.exec_module(module)
    return module


build_city = load_builder()

COLUMNS = ["Level1", "Level 2", "Level 3", "Level4", "Code", "Active CBP",
           "Draft CBP", "Total CBP", "CBP used", "% AVA Sourcing",
           "% Ariba Sourcing", "Spend FY26/27", "AI generated RFPs",
           "MarketsList", "VPC Category Manager", "Definition"]


def row(level1, l2, l3, l4, code, active=0, draft=0, used=0, ai=0,
        spend=0, markets="", manager="", definition=""):
    return [level1, l2, l3, l4, code, active, draft, active + draft, used,
            1, 0, spend, ai, markets, manager, definition]


def workbook_of(sheets: dict) -> pathlib.Path:
    """A workbook in the shape the estates were asked to supply."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        sheet.append(COLUMNS)
        for r in rows:
            sheet.append(r)
    path = pathlib.Path(tempfile.mkdtemp()) / "extract.xlsx"
    book.save(path)
    return path


ALPHA = [
    row("Alpha", "Radio", "Access", "Small Cells", "A101", active=1, draft=1,
        used=2, ai=2, spend=40_000_000, markets="Germany, Italy",
        manager="A Person", definition="Radio access."),
    row("Alpha", "Radio", "Access", "Antennas", "A102", active=1,
        spend=5_000_000, markets="Germany", manager="A Person"),
    row("Alpha", "Core", "Switching", "Packet", "A201", spend=1_000_000),
]
BETA = [
    row("Beta", "Offices", "Leasing", "Leases", "B101", active=1, draft=2,
        used=1, spend=9_000_000, markets="Spain"),
    row("Beta", "Offices", "Leasing", "Fit-out", "B102", draft=1, spend=2_000_000),
]


class EstateSelection(unittest.TestCase):
    """Which estates the build reads, and in what order."""

    def test_an_absent_block_still_builds_the_single_estate(self):
        """A config from before the estates existed has to keep working."""
        only = build_city.estates_in({})
        self.assertEqual(1, len(only))
        self.assertEqual(build_city.SCOPE_L1, only[0]["level1"])

    def test_a_disabled_estate_is_not_read(self):
        picked = build_city.estates_in({"estates": {"list": [
            {"id": "a", "label": "A", "level1": "Alpha", "enabled": True},
            {"id": "b", "label": "B", "level1": "Beta", "enabled": False},
        ]}})
        self.assertEqual(["a"], [e["id"] for e in picked])

    def test_estates_come_back_in_their_configured_order(self):
        picked = build_city.estates_in({"estates": {"list": [
            {"id": "c", "label": "C", "level1": "C", "order": 2},
            {"id": "a", "label": "A", "level1": "A", "order": 0},
            {"id": "b", "label": "B", "level1": "B", "order": 1},
        ]}})
        self.assertEqual(["a", "b", "c"], [e["id"] for e in picked])

    def test_what_is_switched_on_is_what_was_built(self):
        """The config and the committed data have to agree.

        Switching an estate on without rebuilding leaves an island configured
        and absent, and rebuilding without switching one on leaves the reverse.
        Neither shows up anywhere else: the renderer draws whatever is in the
        file and the config is only read at build time.
        """
        import json
        config = build_city.load_config()
        live = [e["id"] for e in build_city.estates_in(config)]
        built = [e["id"] for e in json.loads(
            (REPO / "data" / "city.json").read_text(encoding="utf-8"))["estates"]]
        self.assertEqual(sorted(live), sorted(built),
                         "config/metrics.yaml and data/city.json disagree about "
                         "which estates exist. Run: python3 run.py build")
        listed = {e["id"] for e in config["estates"]["list"]}
        self.assertEqual({"network", "itande", "commercial", "corporate"}, listed,
                         "all four estates should stay configured")


class ReadingAnEstate(unittest.TestCase):
    """One estate's categories, read from its own sheet."""

    @classmethod
    def setUpClass(cls):
        cls.path = workbook_of({"Alpha": ALPHA, "Beta": BETA})
        cls.book = openpyxl.load_workbook(cls.path, read_only=True, data_only=True)
        cls.config = build_city.load_config()

    def read(self, estate):
        return build_city.read_estate(self.book, None, estate, self.config)

    def test_an_estate_reads_only_its_own_sheet(self):
        alpha = self.read({"id": "a", "label": "A", "level1": "Alpha", "sheet": "Alpha"})
        self.assertEqual(["A101", "A102", "A201"], sorted(c["code"] for c in alpha))

    def test_every_category_carries_the_island_it_stands_on(self):
        beta = self.read({"id": "b", "label": "B", "level1": "Beta", "sheet": "Beta"})
        self.assertEqual({"b"}, {c["estate"] for c in beta})

    def test_markets_are_read_from_the_estate_sheet(self):
        alpha = self.read({"id": "a", "label": "A", "level1": "Alpha", "sheet": "Alpha"})
        found = {c["code"]: c["markets"] for c in alpha}
        self.assertEqual(["Germany", "Italy"], found["A101"])
        self.assertEqual([], found["A201"], "a lot with no blueprint has no markets")

    def test_the_journey_score_does_not_depend_on_the_market_list(self):
        """Why a missing MarketsList does not flatten a new estate.

        The blueprint rung is decided by the blueprint counts. A201 has none,
        A102 has one, A101 has two, and that is the whole ladder.
        """
        alpha = {c["code"]: c for c in
                 self.read({"id": "a", "label": "A", "level1": "Alpha", "sheet": "Alpha"})}
        self.assertEqual(0, alpha["A201"]["journey"]["blueprint"])
        self.assertEqual(25, alpha["A102"]["journey"]["blueprint"])
        self.assertEqual(40, alpha["A101"]["journey"]["blueprint"])

    def test_an_estate_with_no_rows_reads_as_empty_rather_than_failing(self):
        self.assertEqual([], self.read(
            {"id": "z", "label": "Z", "level1": "Zeta", "sheet": "Alpha"}))

    def test_districts_are_tagged_and_rolled_up_per_island(self):
        alpha = self.read({"id": "a", "label": "A", "level1": "Alpha", "sheet": "Alpha"})
        districts = build_city.districts_of(alpha, "a", self.config)
        self.assertEqual(["Core", "Radio"], [d["name"] for d in districts])
        self.assertEqual({"a"}, {d["estate"] for d in districts})
        core = next(d for d in districts if d["name"] == "Core")
        self.assertEqual(0, core["totals"]["journey"]["total"],
                         "a district of one empty lot scores nothing")


class AddingAnEstateChangesNothingElse(unittest.TestCase):
    """The promise that matters most while the new estates are arriving."""

    @classmethod
    def setUpClass(cls):
        cls.path = workbook_of({"Alpha": ALPHA, "Beta": BETA})
        cls.book = openpyxl.load_workbook(cls.path, read_only=True, data_only=True)
        cls.config = build_city.load_config()
        cls.alpha = {"id": "a", "label": "A", "level1": "Alpha", "sheet": "Alpha"}
        cls.beta = {"id": "b", "label": "B", "level1": "Beta", "sheet": "Beta"}

    def test_reading_a_second_estate_does_not_disturb_the_first(self):
        alone = build_city.read_estate(self.book, None, self.alpha, self.config)
        build_city.read_estate(self.book, None, self.beta, self.config)
        again = build_city.read_estate(self.book, None, self.alpha, self.config)
        self.assertEqual(alone, again)

    def test_codes_do_not_collide_between_estates(self):
        """The join key across the whole region.

        Nothing in the taxonomy guarantees this, so it is asserted rather than
        assumed. A collision would silently merge two categories.
        """
        alpha = build_city.read_estate(self.book, None, self.alpha, self.config)
        beta = build_city.read_estate(self.book, None, self.beta, self.config)
        overlap = {c["code"] for c in alpha} & {c["code"] for c in beta}
        self.assertEqual(set(), overlap)


class MonumentsAcrossTheRegion(unittest.TestCase):
    """One monument stands once on the sea, or not at all.

    A monument says which market carried a blueprint furthest. The same one
    on two islands says that twice about two different categories, which is
    not a thing the data can mean, and a viewer who recognises the shape has
    no way to tell which of the two they are looking at.

    Only Networks has market lists today, so nothing collides in the shipped
    file. These fixtures give two estates the same market so the collision is
    forced rather than waited for.
    """

    SAME_MARKET = {
        "One": [
            row("One", "Radio", "Access", "Alpha", "A101", active=2, draft=0,
                used=2, ai=2, spend=40_000_000, markets="UK"),
            row("One", "Radio", "Access", "Beta", "A102", active=2, draft=0,
                used=2, ai=2, spend=30_000_000, markets="UK"),
        ],
        # The second island reaches the same market first and a second one
        # after it, so there is somewhere for it to fall through to.
        "Two": [
            row("Two", "Offices", "Leasing", "Gamma", "B101", active=2, draft=0,
                used=2, ai=2, spend=40_000_000, markets="UK, Italy"),
            row("Two", "Offices", "Leasing", "Delta", "B102", active=2, draft=0,
                used=2, ai=2, spend=30_000_000, markets="UK, Italy"),
        ],
    }

    @classmethod
    def setUpClass(cls):
        cls.path = workbook_of(cls.SAME_MARKET)
        cls.book = openpyxl.load_workbook(cls.path, read_only=True, data_only=True)
        cls.config = build_city.load_config()
        cls.one = {"id": "one", "label": "One", "level1": "One", "sheet": "One"}
        cls.two = {"id": "two", "label": "Two", "level1": "Two", "sheet": "Two"}

    def named(self, categories):
        return [c["landmark"]["name"] for c in categories if c.get("landmark")]

    def test_islands_sharing_a_market_do_not_share_a_monument(self):
        claimed = set()
        first = build_city.read_estate(
            self.book, None, self.one, self.config, claimed)
        second = build_city.read_estate(
            self.book, None, self.two, self.config, claimed)
        awarded = self.named(first) + self.named(second)
        self.assertEqual(len(awarded), len(set(awarded)),
                         f"a monument was awarded twice: {awarded}")

    def test_without_the_shared_claim_set_they_would(self):
        """What the shared set is for, asserted so it cannot be dropped.

        Each island claiming independently is the behaviour this replaced. If
        this ever stops reproducing, the two cases have converged and the
        check above is no longer testing anything.
        """
        first = build_city.read_estate(self.book, None, self.one, self.config)
        second = build_city.read_estate(self.book, None, self.two, self.config)
        awarded = self.named(first) + self.named(second)
        self.assertLess(len(set(awarded)), len(awarded))

    def test_a_claimed_monument_sends_the_island_to_its_next_market(self):
        """A taken monument costs that market, not the category.

        The category has more than one market behind it, so the scarce
        monument should move the award along rather than cancel it. Stopping
        at the first market would cost a qualifying category its landmark for
        a reason that has nothing to do with it.
        """
        claimed = set()
        first = build_city.read_estate(
            self.book, None, self.one, self.config, claimed)
        second = build_city.read_estate(
            self.book, None, self.two, self.config, claimed)
        self.assertIn("Big Ben", self.named(first))
        self.assertTrue(self.named(second),
                        "the second island qualified and got nothing")
        self.assertNotIn("Big Ben", self.named(second))


class TheShippedData(unittest.TestCase):
    """What data/city.json actually carries."""

    @classmethod
    def setUpClass(cls):
        import json
        cls.city = json.loads((REPO / "data" / "city.json").read_text(encoding="utf-8"))

    def test_the_estate_roster_matches_the_categories(self):
        listed = {e["id"] for e in self.city["estates"]}
        present = {c["estate"] for c in self.city["categories"]}
        self.assertEqual(listed, present)

    def test_every_estate_counts_what_it_holds(self):
        for estate in self.city["estates"]:
            mine = [c for c in self.city["categories"] if c["estate"] == estate["id"]]
            self.assertEqual(len(mine), estate["counts"]["categories"], estate["id"])
            self.assertEqual(
                len({c["district"] for c in mine}), estate["counts"]["districts"],
                estate["id"])

    def test_one_estate_rolls_up_to_the_region(self):
        """With a single estate the two figures are the same number.

        It stops being a tautology the moment a second estate is switched on,
        which is exactly when a mistake here would be hardest to see.
        """
        if len(self.city["estates"]) != 1:
            self.skipTest("more than one estate is enabled")
        self.assertEqual(self.city["totals"], self.city["estates"][0]["totals"])

    def test_every_district_belongs_to_an_estate_that_exists(self):
        listed = {e["id"] for e in self.city["estates"]}
        for district in self.city["districts"]:
            self.assertIn(district["estate"], listed, district["name"])


if __name__ == "__main__":
    unittest.main()
